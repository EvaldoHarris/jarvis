import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from core import document_workflow as docs
from actions.document_workflow import document_workflow
from core.action_loader import ActionRegistry, _validate
from actions import document_workflow as action
from core import task_plan


def text_pdf(path, text=None):
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    page = writer.pages[0]
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 30 750 Td ({text}) Tj ET'.encode('latin-1'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as output:
        writer.write(output)


class DocumentWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def text(self, name, content):
        path = self.root / name
        path.write_text(content, encoding='utf-8')
        return path

    def test_content_search_matches_accents_and_retains_reference(self):
        path = self.text('proposta.txt', 'Fornecedor A\nOrçamento para revisão\nTotal: R$ 1.234,56')
        result = docs.search_documents(str(self.root), query='ORCAMENTO')
        self.assertEqual(result['matches'][0]['path'], str(path))
        self.assertEqual(result['matches'][0]['reference'], 'linha 2')
        self.assertIn('Orçamento', result['matches'][0]['excerpt'])

    def test_search_is_read_only_skips_private_dependency_folders_and_respects_limit(self):
        self.text('one.txt', 'needle')
        self.text('two.txt', 'needle')
        hidden = self.root / '.git'
        hidden.mkdir()
        (hidden / 'secret.txt').write_text('needle')
        result = docs.search_documents(str(self.root), query='needle', max_results=1)
        self.assertEqual(len(result['matches']), 1)
        self.assertTrue(any('limit' in warning for warning in result['warnings']))
        result = docs.search_documents(str(self.root), query='needle')
        self.assertEqual(len(result['matches']), 2)

    def test_pdf_amount_has_page_reference(self):
        path = self.root / 'budget.pdf'
        text_pdf(path, 'Total: R$ 1.234,56')
        doc = docs.read_document(path)
        self.assertIn('R$ 1.234,56', doc['segments'][0]['text'])
        self.assertEqual(doc['segments'][0]['reference'], 'pagina 1')

    def test_image_or_blank_pdf_warns_about_ocr_instead_of_inventing_price(self):
        path = self.root / 'scanned.pdf'
        text_pdf(path)
        doc = docs.read_document(path)
        self.assertEqual(doc['segments'], [])
        self.assertTrue(any('OCR' in warning for warning in doc['warnings']))

    def test_docx_reads_paragraphs_and_table_with_source_locations(self):
        from docx import Document
        path = self.root / 'proposal.docx'
        document = Document()
        document.add_paragraph('Fornecedor A')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Total'
        table.cell(0, 1).text = 'R$ 250,00'
        document.save(path)
        doc = docs.read_document(path)
        self.assertEqual(doc['segments'][0]['reference'], 'paragrafo 1')
        self.assertIn('tabela 1, linha 1', doc['segments'][1]['reference'])

    def test_read_pagination_rejects_mixed_document_revisions(self):
        path = self.text('long.txt', 'alpha\nbeta\ngamma')
        digest = docs.read_document(path)['sha256']
        first = document_workflow({'action': 'read', 'path': str(path), 'limit': 1})
        self.assertIn('"next_offset": 1', first)
        path.write_text('changed', encoding='utf-8')
        second = document_workflow({'action': 'read', 'path': str(path), 'offset': 1, 'expected_sha256': digest})
        self.assertIn('status=failed', second)
        self.assertIn('changed', second)

    def test_compare_keeps_total_freight_and_installments_separate(self):
        a = self.text('A.txt', 'Total: R$ 1.234,56\nFrete: R$ 50,00\nParcela: R$ 102,88')
        b = self.text('B.txt', 'Total: R$ 900,00')
        comparison = docs.compare_amounts([str(a), str(b)])
        self.assertEqual([row['amount'] for row in comparison['rows']], ['1234.56', '50.00', '102.88', '900.00'])
        self.assertEqual([row['label'] for row in comparison['rows']], ['Total', 'Frete', 'Parcela', 'Total'])
        self.assertTrue(all(row['reference'].startswith('linha ') for row in comparison['rows']))
        self.assertEqual(comparison['errors'], [])

    def test_bad_numeric_format_does_not_become_partial_invented_amount(self):
        a = self.text('A.txt', 'Total R$ 1,2345')
        b = self.text('B.txt', 'Total R$ 300,00')
        comparison = docs.compare_amounts([str(a), str(b)])
        self.assertEqual(len(comparison['rows']), 1)
        self.assertEqual(comparison['errors'][0]['path'], str(a))

    def test_amount_split_at_pagination_boundary_is_reassembled_for_comparison(self):
        a = self.text('A.txt', ('x' * 2997) + 'R$ 1.234,56')
        b = self.text('B.txt', 'R$ 300,00')
        comparison = docs.compare_amounts([str(a), str(b)])
        self.assertEqual(comparison['rows'][0]['amount'], '1234.56')
        self.assertIn('R$ 1.234,56', comparison['rows'][0]['excerpt'])

    def test_real_pdf_and_docx_compare_and_export_with_both_source_references(self):
        from docx import Document
        pdf = self.root / 'supplier_A.pdf'
        word = self.root / 'supplier_B.docx'
        text_pdf(pdf, 'Total: R$ 1.234,56')
        document = Document()
        document.add_paragraph('Total: R$ 900,00')
        document.save(word)
        output = self.root / 'mixed-comparison.csv'
        data = docs.compare_amounts([str(pdf), str(word)])
        self.assertEqual(data['errors'], [])
        saved = docs.save_comparison(data, str(output))
        self.assertTrue(saved['verified'])
        rows = list(csv.DictReader(io.StringIO(output.read_text(encoding='utf-8-sig'))))
        self.assertEqual([row['amount'] for row in rows], ['1234.56', '900.00'])
        self.assertEqual([row['reference'] for row in rows], ['pagina 1', 'paragrafo 1'])
        self.assertIn('DOCX extraction', rows[1]['document_warnings'])

    def test_reading_limits_and_output_undo_are_explicit(self):
        a = self.text('A.txt', 'Total R$ 100,00')
        b = self.text('B.txt', 'Total R$ 200,00')
        with patch.object(docs, 'MAX_BYTES', 2):
            with self.assertRaises(ValueError):
                docs.read_document(a)
        output = self.root / 'comparison.csv'
        with patch('core.undo.push_undo') as undo:
            docs.save_comparison(docs.compare_amounts([str(a), str(b)]), str(output))
        self.assertIn('removed', undo.call_args.args[1]())
        self.assertFalse(output.exists())
        self.assertTrue(a.exists())
        self.assertTrue(b.exists())

    def test_missing_or_duplicate_documents_are_reported_and_do_not_export(self):
        a = self.text('A.txt', 'R$ 100,00')
        output = self.root / 'comparison.csv'
        reply = document_workflow({'action': 'compare', 'paths': [str(a), str(self.root / 'missing.pdf')], 'destination': str(output)})
        self.assertIn('status=unverified', reply)
        self.assertIn('CSV not created', reply)
        self.assertFalse(output.exists())
        duplicate = docs.compare_amounts([str(a), str(a)])
        self.assertTrue(duplicate['errors'])

    def test_csv_readback_and_undo_preserve_sources_and_edited_output(self):
        a = self.text('A.txt', 'Total R$ 100,00')
        b = self.text('B.txt', 'Total R$ 200,00')
        output = self.root / 'comparison.csv'
        data = docs.compare_amounts([str(a), str(b)])
        with patch('core.undo.push_undo') as undo:
            saved = docs.save_comparison(data, str(output))
        self.assertTrue(saved['verified'])
        rows = list(csv.DictReader(io.StringIO(output.read_text(encoding='utf-8-sig'))))
        self.assertEqual(rows[0]['amount'], '100.00')
        self.assertEqual(rows[0]['source'], str(a))
        self.assertEqual(a.read_text(encoding='utf-8'), 'Total R$ 100,00')
        output.write_text('user edit')
        self.assertIn('preserved', undo.call_args.args[1]())
        self.assertTrue(output.exists())

    def test_existing_output_is_never_overwritten(self):
        a = self.text('A.txt', 'R$ 100,00')
        b = self.text('B.txt', 'R$ 200,00')
        output = self.text('comparison.csv', 'existing user data')
        with self.assertRaises(FileExistsError):
            docs.save_comparison(docs.compare_amounts([str(a), str(b)]), str(output))
        self.assertEqual(output.read_text(encoding='utf-8'), 'existing user data')

    def test_formula_like_document_labels_are_escaped_in_csv(self):
        a = self.text('A.txt', '=HYPERLINK("bad") R$ 100,00')
        b = self.text('B.txt', '@formula R$ 200,00')
        output = self.root / 'comparison.csv'
        docs.save_comparison(docs.compare_amounts([str(a), str(b)]), str(output))
        rows = list(csv.DictReader(io.StringIO(output.read_text(encoding='utf-8-sig'))))
        self.assertTrue(rows[0]['label'].startswith("'="))
        self.assertTrue(rows[1]['label'].startswith("'@"))

    def test_end_to_end_registry_find_read_compare_save_and_checkpoint(self):
        a = self.text('budget_A.txt', 'Total R$ 100,00')
        b = self.text('budget_B.txt', 'Total R$ 200,00')
        plans = task_plan.TaskPlans(self.root / 'plans.db')
        plan = plans.update({'action': 'create', 'objective': 'Compare budgets', 'steps': ['Find budgets', 'Read A', 'Save comparison']})
        record = _validate(action, 'document_workflow.py')
        registry = ActionRegistry({'document_workflow': record}, logger=lambda message: None)
        calls = [
            {'action': 'search', 'path': str(self.root), 'name': 'budget'},
            {'action': 'read', 'path': str(a)},
            {'action': 'compare', 'paths': [str(a), str(b)], 'destination': str(self.root / 'comparison.csv')},
        ]
        import re
        with patch.object(task_plan, '_path', return_value=plans.path):
            for params in calls:
                reply = registry.run('document_workflow', params)
                receipt = re.search(r'task_receipt=([a-f0-9]+)', reply)[1]
                plan = plans.update({'action': 'complete_step', 'plan_id': plan['id'], 'revision': plan['revision'],
                                     'receipt_id': receipt, 'evidence': f"Observed document workflow: {params['action']}"})
        self.assertEqual(plan['state'], 'completed')
        self.assertTrue((self.root / 'comparison.csv').exists())


if __name__ == '__main__':
    unittest.main()
