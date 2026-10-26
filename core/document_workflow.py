"""Bounded local document discovery, sourced monetary extraction and CSV output."""
import csv
import hashlib
import io
import os
import re
import time
import unicodedata
from decimal import Decimal
from pathlib import Path

SUPPORTED = {'.txt', '.md', '.csv', '.pdf', '.docx'}
MAX_BYTES = 10_000_000
MAX_CHARS = 200_000
SKIP_DIRS = {'.git', '.venv', 'venv', 'node_modules', '__pycache__', '.idea'}
MONEY = re.compile(r'R\$\s*(-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{2})?)(?![\d.,])')


def fold(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', value.casefold()) if not unicodedata.combining(c))


def source_path(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Provide an explicit document or folder path')
    return Path(value).expanduser().resolve(strict=True)


def read_document(value, use_ocr=False):
    path = source_path(str(value))
    if not path.is_file() or path.suffix.lower() not in SUPPORTED:
        raise ValueError('Supported documents: PDF, DOCX, TXT, MD, CSV')
    if path.stat().st_size > MAX_BYTES:
        raise ValueError('Document exceeds the 10 MB reading limit')
    with path.open('rb') as stream:
        raw = stream.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValueError('Document changed or exceeds the 10 MB reading limit')
    digest = hashlib.sha256(raw).hexdigest()
    segments, warnings = [], []
    size = 0

    def add(reference, text):
        nonlocal size
        text = str(text or '')
        if size >= MAX_CHARS:
            return
        if len(text) > MAX_CHARS - size:
            warnings.append('Text extraction truncated at 200000 characters; unread content is not covered')
        text = text[:MAX_CHARS - size]
        size += len(text)
        # Keep independently addressable chunks for pagination and citations.
        for offset in range(0, len(text), 3000):
            segments.append({'reference': reference, 'part': offset // 3000 + 1, 'text': text[offset:offset + 3000]})

    ext = path.suffix.lower()
    if ext == '.pdf':
        import pdfplumber
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            if len(pdf.pages) > 100:
                warnings.append('PDF limited to the first 100 pages')
            for index, page in enumerate(pdf.pages[:100]):
                if size >= MAX_CHARS:
                    warnings.append('PDF remaining pages were not extracted because the text limit was reached')
                    break
                text = page.extract_text() or ''
                if not text.strip():
                    if use_ocr and index < 10:
                        from core.local_ocr import read_ocr
                        extracted = read_ocr(str(path), first_page=index+1, count=1)
                        text = extracted['segments'][0]['text']
                        warnings.append(f'Page {index + 1}: OCR transcription, verify monetary amounts against the source image')
                    else:
                        warnings.append(f'Page {index + 1}: no selectable text; may need OCR, not proof of an empty page')
                add(f'pagina {index + 1}', text)
    elif ext == '.docx':
        import zipfile
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                raise ValueError('Expanded DOCX exceeds the 50 MB limit')
        document = Document(io.BytesIO(raw))
        for index, paragraph in enumerate(document.paragraphs):
            add(f'paragrafo {index + 1}', paragraph.text)
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                add(f'tabela {table_index + 1}, linha {row_index + 1}', ' | '.join(cell.text for cell in row.cells))
        warnings.append('DOCX extraction covers body paragraphs/tables; headers, text boxes and embedded images are not covered')
    else:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1252')
            warnings.append('Text decoded as Windows-1252 because UTF-8 decoding failed')
        for index, line in enumerate(text.splitlines()):
            add(f'linha {index + 1}', line)
            if size >= MAX_CHARS:
                warnings.append('Remaining lines were not read because the text limit was reached')
                break
    if not segments:
        warnings.append('No text extracted; do not infer that the document has no relevant information')
    return {'path': str(path), 'sha256': digest, 'segments': segments, 'warnings': list(dict.fromkeys(warnings))}


def search_documents(folder, name='', query='', max_results=20):
    root = source_path(folder)
    if not root.is_dir():
        raise ValueError('Search requires a folder')
    if not name.strip() and not query.strip():
        raise ValueError('Provide part of a filename or a content query')
    results, warnings = [], []
    inspected, directories = 0, 0
    deadline = time.monotonic() + 20
    def walk_error(error):
        warnings.append(f'Folder unavailable: {error.filename}')
    for parent, dirs, files in os.walk(root, followlinks=False, onerror=walk_error):
        directories += 1
        dirs[:] = sorted(d for d in dirs if d not in SKIP_DIRS and not d.startswith('.')
                         and not (Path(parent) / d).is_symlink()
                         and (Path(parent) / d).resolve().is_relative_to(root))
        if directories > 500 or time.monotonic() > deadline:
            warnings.append('Search stopped at the folder/time limit; results are partial')
            break
        for filename in sorted(files):
            path = Path(parent) / filename
            if path.suffix.lower() not in SUPPORTED or path.is_symlink():
                continue
            if inspected >= 200 or time.monotonic() > deadline:
                warnings.append('Search stopped at the 200-document/time limit; results are partial')
                return {'matches': results, 'scanned': inspected, 'warnings': list(dict.fromkeys(warnings))}
            inspected += 1
            if name and fold(name) not in fold(filename):
                continue
            if not path.resolve().is_relative_to(root):
                continue
            try:
                if query:
                    doc = read_document(path)
                    hits = [segment for segment in doc['segments'] if fold(query) in fold(segment['text'])]
                    warnings.extend(f'{filename}: {warning}' for warning in doc['warnings'])
                    if not hits:
                        continue
                    snippet = hits[0]
                    text = snippet['text']
                    position = fold(text).find(fold(query))
                    start = max(0, position - 150)
                    results.append({'path': str(path), 'reference': snippet['reference'], 'excerpt': text[start:start + 500], 'sha256': doc['sha256']})
                else:
                    results.append({'path': str(path), 'matched': 'filename_only'})
            except Exception as exc:
                warnings.append(f'{filename}: {exc}')
            if len(results) >= max_results:
                warnings.append('Result limit reached; additional matches may exist')
                return {'matches': results, 'scanned': inspected, 'warnings': list(dict.fromkeys(warnings))}
    return {'matches': results, 'scanned': inspected, 'warnings': list(dict.fromkeys(warnings))}


def compare_amounts(paths, use_ocr=False):
    if not isinstance(paths, list) or not 2 <= len(paths) <= 10:
        raise ValueError('Select 2 to 10 documents explicitly; do not compare arbitrary search matches')
    rows, documents, errors = [], [], []
    resolved = set()
    for path in paths:
        if len(rows) >= 500:
            errors.append({'path': str(path), 'reason': 'Monetary extraction limit reached; remaining documents not compared'})
            continue
        try:
            doc = read_document(path, use_ocr=use_ocr)
            if doc['path'] in resolved:
                raise ValueError('The same document was selected more than once')
            resolved.add(doc['path'])
            documents.append({key: doc[key] for key in ('path', 'sha256', 'warnings')})
            count = 0
            references = {}
            for segment in doc['segments']:
                references[segment['reference']] = references.get(segment['reference'], '') + segment['text']
            for reference, text in references.items():
                for line in text.splitlines():
                    for match in MONEY.finditer(line):
                        if len(rows) >= 500:
                            raise ValueError('500 monetary values reached; comparison is partial')
                        amount = Decimal(match[1].replace('.', '').replace(',', '.'))
                        prefix = line[:match.start()]
                        label = prefix.strip(' :;|\t')[-100:] or 'valor sem rotulo'
                        rows.append({'source': doc['path'], 'sha256': doc['sha256'], 'reference': reference,
                                     'label': label, 'currency': 'BRL', 'amount': format(amount, '.2f'),
                                     'excerpt': line[max(0, match.start() - 150):match.end() + 200]})
                        count += 1
            if not count:
                errors.append({'path': doc['path'], 'reason': 'No explicit R$ values in Brazilian numeric format were extracted; no price inferred'})
        except Exception as exc:
            errors.append({'path': str(path), 'reason': str(exc)})
    return {'documents': documents, 'rows': rows, 'errors': errors,
            'interpretation': 'Values are source observations, not an automatic cheapest-offer recommendation. Totals, installments, freight, quantities and conditions may differ. No currency conversion or model inference was performed.'}


def save_comparison(data, destination):
    from core.undo import push_undo
    if data['errors'] or not data['rows']:
        raise ValueError('Comparison has unread/unpriced documents or incomplete extraction; no CSV created. Review errors instead of silently omitting documents.')
    if not isinstance(destination, str) or not destination.strip():
        raise ValueError('Provide an explicit CSV destination')
    target = Path(destination).expanduser().resolve()
    if target.suffix.lower() == '.xlsx':
        from core.workbooks import create
        fields = ['source', 'reference', 'label', 'currency', 'amount', 'excerpt', 'sha256', 'document_warnings']
        rows = []
        for row in data['rows']:
            notes = next((doc['warnings'] for doc in data['documents'] if doc['path'] == row['source']), [])
            values = {**row, 'amount': float(row['amount']), 'document_warnings': ' | '.join(notes)}
            rows.append([values[key] for key in fields])
        return {**create(fields, rows, str(target), sheet='Comparacao'), 'rows': len(rows)}
    if target.suffix.lower() != '.csv' or not target.parent.is_dir():
        raise ValueError('Destination must be a .csv file in an existing folder')
    fields = ['source', 'reference', 'label', 'currency', 'amount', 'excerpt', 'sha256', 'document_warnings']
    buffer = io.StringIO(newline='')
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    def cell(value):
        text = str(value)
        return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) else text
    for row in data['rows']:
        notes = next((doc['warnings'] for doc in data['documents'] if doc['path'] == row['source']), [])
        values = {**row, 'document_warnings': ' | '.join(notes)}
        writer.writerow({key: cell(values[key]) for key in fields})
    content = buffer.getvalue().encode('utf-8-sig')
    with target.open('xb') as stream:
        stream.write(content)
    digest = hashlib.sha256(content).hexdigest()
    def undo():
        if not target.exists():
            return 'Comparison file already absent'
        if target.stat().st_size != len(content):
            return 'Comparison was edited after creation; file preserved'
        if hashlib.sha256(target.read_bytes()).hexdigest() != digest:
            return 'Comparison was edited after creation; file preserved'
        target.unlink()
        return 'Created comparison file removed; source documents unchanged'
    push_undo(f'criou comparacao {target.name}', undo)
    if target.read_bytes() != content:
        raise RuntimeError('Output was created but read-back verification failed; inspect before retrying')
    return {'path': str(target), 'sha256': digest, 'rows': len(data['rows']), 'verified': True}
