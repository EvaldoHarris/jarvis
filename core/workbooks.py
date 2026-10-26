import hashlib
import io
import math
from pathlib import Path


def value_ok(value):
    if value is not None and (type(value) not in (str, int, float, bool) or
            isinstance(value, str) and len(value) > 30000 or isinstance(value, float) and not math.isfinite(value)):
        raise ValueError('Use bounded plain cell values; no objects or nonfinite numbers')


def write_cell(cell, value):
    value_ok(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # text beginning '=' is never promoted to a formula


def persist(workbook, destination):
    from openpyxl import load_workbook
    from core.undo import push_undo
    target = Path(destination).expanduser().resolve()
    if target.suffix.lower() != '.xlsx' or not target.parent.is_dir():
        raise ValueError('Choose a new .xlsx file in an existing folder')
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    with target.open('xb') as stream:
        stream.write(content)
    with target.open('rb') as stream:
        if stream.read() != content:
            raise RuntimeError('Workbook created but byte verification failed; inspect before retrying')
    reopened = load_workbook(target, read_only=True, data_only=False)
    sheets = reopened.sheetnames
    reopened.close()
    digest = hashlib.sha256(content).hexdigest()
    def undo():
        if not target.exists():
            return 'Generated workbook already absent'
        if target.stat().st_size != len(content) or hashlib.sha256(target.read_bytes()).hexdigest() != digest:
            return 'Workbook changed after creation; preserved'
        target.unlink()
        return 'Generated workbook removed; sources unchanged'
    push_undo('criou planilha ' + target.name, undo)
    return {'path': str(target), 'sha256': digest, 'sheets': sheets, 'verified': True}


def create(headers, rows, destination, sheet='Dados'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if not isinstance(headers, list) or not 1 <= len(headers) <= 40 or any(not isinstance(h, str) or not h for h in headers):
        raise ValueError('Provide 1-40 nonempty column names')
    if not isinstance(rows, list) or len(rows) > 1000 or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise ValueError('Provide up to 1000 rows matching the column count')
    workbook = Workbook()
    page = workbook.active
    page.title = sheet
    for index, row in enumerate([headers, *rows], start=1):
        for column, value in enumerate(row, start=1):
            write_cell(page.cell(index, column), value)
    for cell in page[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='123B4B')
    page.freeze_panes = 'A2'
    page.auto_filter.ref = page.dimensions
    for column, name in enumerate(headers, start=1):
        page.column_dimensions[get_column_letter(column)].width = min(50, max(16, len(name)+3))
    return persist(workbook, destination)


def read_source(path):
    import zipfile
    source = Path(path).expanduser().resolve(strict=True)
    if source.suffix.lower() != '.xlsx' or source.stat().st_size > 10_000_000:
        raise ValueError('Choose an XLSX workbook no larger than 10 MB')
    raw = source.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(row.file_size for row in archive.infolist()) > 50_000_000:
            raise ValueError('Expanded workbook exceeds 50 MB')
    return source, raw, hashlib.sha256(raw).hexdigest()


def inspect(path, sheet=None, first_row=1, count=30):
    from openpyxl import load_workbook
    source, raw, digest = read_source(path)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        page = workbook[sheet] if sheet else workbook.active
        if not 1 <= first_row or not 1 <= count <= 100:
            raise ValueError('Read 1-100 rows from an explicit positive row index')
        rows = [[str(v) if v is not None and not isinstance(v,(str,int,float,bool)) else v for v in row]
                for row in page.iter_rows(min_row=first_row, max_row=min(page.max_row, first_row+count-1), max_col=min(page.max_column,40), values_only=True)]
        return {'path': str(source), 'sha256': digest, 'sheets': workbook.sheetnames, 'sheet': page.title,
                'rows': rows, 'first_row': first_row, 'total_rows': page.max_row,
                'note': 'Only requested rows/first 40 columns returned. Formulas are source text, not evaluated results.'}
    finally:
        workbook.close()


def update(path, expected_sha256, cells, destination, sheet=None):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple
    source, raw, digest = read_source(path)
    if not expected_sha256 or digest != expected_sha256:
        raise ValueError('Workbook changed or missing expected_sha256; inspect again')
    if Path(destination).expanduser().resolve() == source:
        raise ValueError('Choose a new destination; original workbook is preserved')
    if not isinstance(cells, dict) or not 1 <= len(cells) <= 200:
        raise ValueError('Provide 1-200 cell addresses and plain values')
    workbook = load_workbook(io.BytesIO(raw), data_only=False)
    page = workbook[sheet] if sheet else workbook.active
    for address, value in cells.items():
        row, column = coordinate_to_tuple(address)
        if not 1 <= row <= 100000 or not 1 <= column <= 16384:
            raise ValueError('Cell address outside supported bounds')
        write_cell(page.cell(row, column), value)
    return persist(workbook, destination)
